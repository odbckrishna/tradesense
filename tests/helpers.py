"""Synthetic market data and OANDA payload builders shared by the test-suite."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[1]

FREQ = {"1m": "1min", "5m": "5min", "15m": "15min"}
GRANULARITY = {"1m": "M1", "5m": "M5", "15m": "M15"}


def make_candles(
    n: int = 300,
    timeframe: str = "5m",
    *,
    seed: int = 7,
    start_price: float = 1.10,
    drift: float = 0.0,
    volatility: float = 0.0006,
    end: str = "2026-03-04 12:00",
    last_incomplete: bool = False,
) -> pd.DataFrame:
    """Random-walk OHLCV candles in the same shape :func:`parse_candles` produces."""
    rng = np.random.default_rng(seed)
    index = pd.date_range(
        end=pd.Timestamp(end, tz="UTC"), periods=n, freq=FREQ[timeframe], name="time"
    )
    returns = rng.normal(drift, volatility, n)
    close = start_price * np.exp(np.cumsum(returns))
    open_ = np.concatenate([[start_price], close[:-1]])
    wick = np.abs(rng.normal(0, volatility * start_price * 0.6, (2, n)))
    high = np.maximum(open_, close) + wick[0]
    low = np.minimum(open_, close) - wick[1]
    volume = rng.integers(20, 400, n).astype(float)
    frame = pd.DataFrame(
        {
            "open": open_,
            "high": high,
            "low": low,
            "close": close,
            "volume": volume,
            "complete": True,
        },
        index=index,
    )
    if last_incomplete:
        frame.iloc[-1, frame.columns.get_loc("complete")] = False
    return frame


def trending_candles(
    direction: str, n: int = 300, timeframe: str = "5m", seed: int = 3
) -> pd.DataFrame:
    drift = 0.0006 if direction == "up" else -0.0006
    return make_candles(n, timeframe, seed=seed, drift=drift, volatility=0.0004)


def oanda_candles_payload(
    df: pd.DataFrame, instrument: str = "EUR_USD", granularity: str = "M5", price_key: str = "mid"
) -> dict[str, Any]:
    candles = []
    for ts, row in df.iterrows():
        candles.append(
            {
                "complete": bool(row["complete"]),
                "volume": int(row["volume"]),
                "time": ts.strftime("%Y-%m-%dT%H:%M:%S.000000000Z"),
                price_key: {
                    "o": f"{row['open']:.5f}",
                    "h": f"{row['high']:.5f}",
                    "l": f"{row['low']:.5f}",
                    "c": f"{row['close']:.5f}",
                },
            }
        )
    return {"instrument": instrument, "granularity": granularity, "candles": candles}


def oanda_bid_ask_payload(
    df: pd.DataFrame, spread: float = 0.00012, granularity: str = "M5"
) -> dict[str, Any]:
    """``price=BA`` candles: bid = the frame's prices, ask = bid + ``spread``."""
    payload = oanda_candles_payload(df, granularity=granularity, price_key="bid")
    for candle in payload["candles"]:
        candle["ask"] = {k: f"{float(v) + spread:.5f}" for k, v in candle["bid"].items()}
    return payload


def oanda_instruments_payload(names: list[str]) -> dict[str, Any]:
    return {
        "instruments": [
            {
                "name": n,
                "type": "CURRENCY",
                "displayName": n.replace("_", "/"),
                "pipLocation": -4,
                "displayPrecision": 5,
            }
            for n in names
        ],
        "lastTransactionID": "1",
    }


# --------------------------------------------------------------------------- sample files


def make_pdf(lines: list[str]) -> bytes:
    """A minimal single-page PDF with real extractable text (Helvetica)."""

    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", r"\(").replace(")", r"\)")

    content = (
        "BT /F1 12 Tf 72 720 Td 16 TL "
        + " ".join(f"({escape(line)}) Tj T*" for line in lines)
        + " ET"
    )
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return out


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import io

    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_image(fmt: str = "PNG", size: tuple[int, int] = (320, 200)) -> bytes:
    """A small 'chart-like' image (two lows + breakout line) in the given format."""
    import io

    from PIL import Image, ImageDraw

    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    w, h = size
    points = [
        (0, h * 0.2),
        (w * 0.25, h * 0.8),
        (w * 0.45, h * 0.45),
        (w * 0.65, h * 0.8),
        (w, h * 0.1),
    ]
    draw.line(points, fill="black", width=3)
    draw.line([(0, h * 0.45), (w, h * 0.45)], fill="blue", width=1)
    buffer = io.BytesIO()
    image.save(buffer, format=fmt)
    return buffer.getvalue()
